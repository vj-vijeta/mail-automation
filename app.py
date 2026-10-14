import streamlit as st
import smtplib
import time
import io
import os
import re
import tempfile
import subprocess
import pandas as pd
from datetime import datetime
import streamlit.components.v1 as components

from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.image import MIMEImage
from email import encoders
from email.utils import formataddr

from streamlit_quill import st_quill
from pptx import Presentation

# --- Page Configuration ---
st.set_page_config(page_title="Ei Mass Emailer", page_icon="✉️", layout="wide")
st.title("✉️ Professional Mass Email Campaigner")
st.write("Run a regular update campaign, or a separate certificate campaign — each with its own audience list and email design.")


# =========================================================================
# HELPER: FLEXIBLE COLUMN NORMALIZATION
# =========================================================================
def normalize_dataframe_columns(df, target_name_col="Name"):
    """
    Standardizes 'email' to 'Email' and the target name to 'Name'. 
    Leaves all other columns (School Name, Event, etc.) intact to be used as dynamic variables.
    """
    if df is None or df.empty:
        return df

    rename_map = {}
    for col in df.columns:
        clean_col = str(col).strip().lower()
        if clean_col in ["email", "email id", "emailid", "mail"]:
            rename_map[col] = "Email"
        elif clean_col == target_name_col.lower():
            rename_map[col] = target_name_col

    df = df.rename(columns=rename_map)
    return df


# =========================================================================
# HELPER: MULTI-VARIABLE ENGINE
# =========================================================================
VAR_RE = re.compile(r"\{\{?\s*([^{}]+?)\s*\}?\}")

def replace_variables(text, row):
    """Replace {Var} / {{Var}} (case/whitespace-insensitive) with values from row."""
    if not text or not row:
        return text
    lookup = {str(k).strip().lower(): v for k, v in row.items()}

    def _sub(m):
        key = m.group(1).strip().lower()
        if key not in lookup:
            return m.group(0)
        val = lookup[key]
        try:
            if pd.isna(val):
                return ""
        except (TypeError, ValueError):
            pass
        return str(val)

    return VAR_RE.sub(_sub, text)

def detect_column(columns, keywords, fallback=0):
    for kw in keywords:
        for i, c in enumerate(columns):
            if kw in str(c).strip().lower():
                return i
    return fallback if columns else 0

def render_variable_chips(df):
    """Show every available variable with the first row's sample value."""
    if df is None or df.empty:
        return
    sample = df.iloc[0].to_dict()
    chips = " ".join(
        f"<span style='display:inline-block;margin:3px;padding:4px 10px;border-radius:14px;"
        f"background:#eef3ff;border:1px solid #c5d5ff;font-size:13px;color:#1f3b7a'>"
        f"<code>{{{c}}}</code> → {str(sample.get(c, ''))[:30]}</span>"
        for c in df.columns
    )
    st.markdown("**🏷️ Available variables** (use in subject, heading, body, button):", unsafe_allow_html=False)
    st.markdown(chips, unsafe_allow_html=True)


# =========================================================================
# CERTIFICATE HELPERS
# =========================================================================
PLACEHOLDER_RE = re.compile(r"\{\{.*?\}\}")

def _iter_text_frames(shapes):
    """Yield every text frame in a shape collection, recursing into groups and tables."""
    for shape in shapes:
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from _iter_text_frames(shape.shapes)
        if shape.has_text_frame:
            yield shape.text_frame
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame

def extract_placeholders(pptx_bytes):
    prs = Presentation(io.BytesIO(pptx_bytes))
    found = set()
    for slide in prs.slides:
        for tf in _iter_text_frames(slide.shapes):
            for para in tf.paragraphs:
                full_text = "".join(r.text for r in para.runs)
                found.update(PLACEHOLDER_RE.findall(full_text))
    return sorted(found)

def fill_certificate(pptx_bytes, mapping):
    prs = Presentation(io.BytesIO(pptx_bytes))
    for slide in prs.slides:
        for tf in _iter_text_frames(slide.shapes):
            for para in tf.paragraphs:
                runs = para.runs
                if not runs:
                    continue
                full_text = "".join(r.text for r in runs)
                if "{{" not in full_text:
                    continue
                new_text = full_text
                for placeholder, value in mapping.items():
                    new_text = new_text.replace(placeholder, str(value))
                if new_text != full_text:
                    runs[0].text = new_text
                    for r in runs[1:]:
                        r.text = ""
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()

def pptx_bytes_to_pdf_bytes(pptx_bytes, timeout=60):
    with tempfile.TemporaryDirectory() as tmp:
        pptx_path = os.path.join(tmp, "cert.pptx")
        with open(pptx_path, "wb") as f:
            f.write(pptx_bytes)
        try:
            subprocess.run(
                ["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, pptx_path],
                check=True, capture_output=True, timeout=timeout,
            )
        except (subprocess.CalledProcessError, FileNotFoundError, subprocess.TimeoutExpired):
            return None
        pdf_path = os.path.join(tmp, "cert.pdf")
        if not os.path.exists(pdf_path):
            return None
        with open(pdf_path, "rb") as f:
            return f.read()

def pptx_bytes_to_image_bytes(pptx_bytes):
    try:
        import fitz  # PyMuPDF
    except ImportError:
        st.error("Missing library. Please add 'PyMuPDF' to your requirements.txt")
        return None
    try:
        pdf_bytes = pptx_bytes_to_pdf_bytes(pptx_bytes)
        if not pdf_bytes:
            return None
            
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        page = doc.load_page(0)
        zoom = fitz.Matrix(2, 2)
        pix = page.get_pixmap(matrix=zoom, alpha=False)
        return pix.tobytes("png")
    except Exception as e:
        st.error(f"Image conversion failed: {e}")
        return None

def sanitize_filename(name):
    name = re.sub(r"[^\w\s-]", "", str(name)).strip()
    return re.sub(r"[-\s]+", "_", name) or "certificate"


# =========================================================================
# REUSABLE UI BLOCKS
# =========================================================================
def render_audience(key_prefix, name_col="Name"):
    tab_csv, tab_manual = st.tabs(["📁 Upload CSV", "✍️ Add Manually"])

    df_csv = None
    with tab_csv:
        uploaded_csv = st.file_uploader(
            f"Upload Data (CSV must include an 'Email' and '{name_col}' column — case insensitive)",
            type="csv", key=f"{key_prefix}_csv_uploader",
        )
        if uploaded_csv:
            try:
                df_csv = pd.read_csv(uploaded_csv)
                df_csv = normalize_dataframe_columns(df_csv, target_name_col=name_col)
            except Exception as e:
                st.error(f"Could not read CSV file: {e}")

    with tab_manual:
        st.caption(
            "Handy for quick tests or small lists. Note: if you need extra per-recipient "
            f"fields (like School Name or Event), use CSV upload instead — manual entry only has {name_col} and Email."
        )
        paste_text = st.text_area(
            f"Paste emails — one per line. Optionally prefix with a name: `{name_col}, email@domain.com`",
            height=100, key=f"{key_prefix}_paste_emails",
        )
        st.write("Or add rows one at a time:")
        manual_df = st.data_editor(
            pd.DataFrame([{name_col: "", "Email": ""}]),
            num_rows="dynamic", use_container_width=True, key=f"{key_prefix}_manual_editor",
        )

    pasted_rows = []
    if paste_text.strip():
        for line in paste_text.strip().splitlines():
            line = line.strip()
            if not line:
                continue
            if "," in line:
                name_part, email_part = line.split(",", 1)
                pasted_rows.append({name_col: name_part.strip(), "Email": email_part.strip()})
            else:
                pasted_rows.append({name_col: "", "Email": line})

    manual_clean = manual_df.copy()
    manual_clean = normalize_dataframe_columns(manual_clean, target_name_col=name_col)
    if "Email" in manual_clean.columns:
        manual_clean["Email"] = manual_clean["Email"].astype(str).str.strip()
        manual_clean = manual_clean[manual_clean["Email"] != ""]

    manual_parts = [p for p in [pd.DataFrame(pasted_rows), manual_clean] if not p.empty]
    manual_combined = pd.concat(manual_parts, ignore_index=True) if manual_parts else pd.DataFrame()

    if df_csv is not None and not df_csv.empty:
        with st.expander("🎛️ Column & Variable Mapping", expanded=True):
            csv_cols = list(df_csv.columns)
            m1, m2 = st.columns(2)
            with m1:
                email_src = st.selectbox(
                    "Which column is the Email address?", csv_cols,
                    index=detect_column(csv_cols, ["email", "mail", "email id", "emailid"]), key=f"{key_prefix}_map_email",
                )
            with m2:
                name_src = st.selectbox(
                    f"Which column is the Recipient Name ({name_col})?", ["(none)"] + csv_cols,
                    index=1 + detect_column(csv_cols, [name_col.lower(), "name", "recipient"]),
                    key=f"{key_prefix}_map_name",
                )

            df_csv = df_csv.copy()
            df_csv["Email"] = df_csv[email_src]
            if name_src != "(none)":
                df_csv[name_col] = df_csv[name_src]

            st.caption(
                "Every CSV column is already a variable, e.g. `{" + str(csv_cols[0]) + "}`. "
                "Optionally add your own variable names and pick which column each one uses:"
            )
            alias_df = st.data_editor(
                pd.DataFrame([{"Variable": "", "CSV Column": None}]),
                num_rows="dynamic", use_container_width=True, key=f"{key_prefix}_alias_editor",
                column_config={
                    "Variable": st.column_config.TextColumn("Variable name (used as {Variable})"),
                    "CSV Column": st.column_config.SelectboxColumn("Takes value from column", options=csv_cols),
                },
            )
            for _, a in alias_df.iterrows():
                var = str(a.get("Variable") or "").strip().strip("{} ")
                src = a.get("CSV Column")
                if var and src in df_csv.columns:
                    df_csv[var] = df_csv[src]

    frames = [f for f in [df_csv, manual_combined] if f is not None and not f.empty]
    df, columns = None, []
    if frames:
        df = pd.concat(frames, ignore_index=True)
        df = df.loc[:, ~df.columns.duplicated()]
        if "Email" in df.columns:
            df["Email"] = df["Email"].astype(str).str.strip()
            df = df.drop_duplicates(subset=["Email"], keep="first").reset_index(drop=True)
        columns = list(df.columns)
        render_variable_chips(df)
        with st.expander("📊 Preview audience data", expanded=False):
            st.metric("Total Recipients", len(df))
            st.dataframe(df, use_container_width=True)

    return df, columns


def _preview_row(key_prefix, df, token):
    if df is not None and len(df) > 0:
        idx = st.number_input(
            f"Preview as recipient # (1–{len(df)})", min_value=1, max_value=len(df), value=1,
            key=f"{key_prefix}_preview_idx_{st.session_state.get(key_prefix + '_pv', 0)}",
        )
        row = df.iloc[int(idx) - 1].to_dict()
        st.caption(f"Previewing: {row.get('Email', '')}")
        return row
    st.caption("Upload an audience to preview with real data.")
    return {token.strip("{}"): "Jane Doe", "School Name": "Sample School", "Event": "Sample Event"}


def render_compose(key_prefix, default_heading, token="{Name}", df=None):
    compose_mode = st.radio(
        "Choose your editor mode:",
        ["🧱 Easy Builder (No Coding)", "💻 Paste Custom HTML Code"],
        horizontal=True, key=f"{key_prefix}_compose_mode",
    )

    generated_html = ""
    local_inline_img = None

    if compose_mode == "🧱 Easy Builder (No Coding)":
        tab1, tab2 = st.tabs(["✍️ Edit Modules", "👁️ Live Preview"])

        with tab1:
            with st.expander("ℹ️ How to use the Easy Builder", expanded=False):
                st.write(f"""
                * **Personalization:** Type `{token}` anywhere to insert the recipient's name. Use `{{Certificate ID}}` to insert their generated ID.
                * **Images:** Upload a local banner, and the app will embed it securely into the email body.
                * **Colors:** Click the color boxes to match your branding.
                """)

            st.subheader("🖼️ Header Banner")
            col1, col2 = st.columns(2)
            with col1:
                local_inline_img = st.file_uploader(
                    "Upload Banner Image (PNG/JPG)", type=['png', 'jpg', 'jpeg'], key=f"{key_prefix}_easy_img"
                )
            with col2:
                header_img_url = st.text_input("OR Paste a Public Image URL", key=f"{key_prefix}_header_img_url")

            img_src = "cid:inline_img" if local_inline_img else header_img_url

            st.subheader("🔤 Heading")
            col3, col4 = st.columns([3, 1])
            with col3:
                heading_text = st.text_input("Heading Text", value=default_heading, key=f"{key_prefix}_heading_text")
            with col4:
                heading_color = st.color_picker("Heading Color", "#0056b3", key=f"{key_prefix}_heading_color")

            st.subheader("📝 Message Body")
            st.caption("Use the toolbar to bold text, add lists, or change font colors.")
            
            # Default HTML Body specifically set to your requirements
            default_body_html = """
            <p>Dear {Name},</p>
            <p>I hope this email finds you well.</p>
            <p>As part of our ongoing outreach to educators and administrators at {School Name}, I am reaching out to introduce you to our upcoming initiative, {Event}.</p>
            <p>Our team has designed this program to support institutions like yours with innovative tools and resources tailored for modern learning environments. We believe this will be a valuable opportunity for your faculty and students to collaborate, explore new methodologies, and enhance their academic experiences.</p>
            <p>We would love to share more details on how your institution can participate. Please feel free to explore the link below to learn more about the agenda and next steps.</p>
            <p>Warm regards,<br><br>Vijeta<br>Educational Initiatives</p>
            """
            
            body_text = st_quill(
                value=default_body_html,
                html=True, key=f"{key_prefix}_quill_builder",
            )

            st.subheader("🔘 Call to Action Button")
            col5, col6, col7 = st.columns([2, 2, 1])
            with col5:
                button_text = st.text_input("Button Text", value="View Portal", key=f"{key_prefix}_button_text")
            with col6:
                button_link = st.text_input("Button Link (URL)", value="https://", key=f"{key_prefix}_button_link")
            with col7:
                button_color = st.color_picker("Button Color", "#0056b3", key=f"{key_prefix}_button_color")

        generated_html = f"""
        <div style="font-family: Arial, sans-serif; color: #333; max-width: 600px; margin: 0 auto; border: 1px solid #ddd; border-radius: 8px; overflow: hidden;">
            {f'<div style="background-color: #f4f6f8; text-align: center; border-bottom: 3px solid {heading_color};"><img src="{img_src}" style="max-width: 100%; height: auto; display: block;"></div>' if img_src else ''}
            <div style="padding: 30px;">
                {f'<h2 style="color: {heading_color}; margin-top: 0;">{heading_text}</h2>' if heading_text else ''}
                <div style="font-size: 15px; line-height: 1.6;">{body_text}</div>
                {f'<div style="text-align: center; margin-top: 30px;"><a href="{button_link}" style="background-color: {button_color}; color: #ffffff; padding: 12px 25px; text-decoration: none; border-radius: 5px; font-weight: bold; display: inline-block;">{button_text}</a></div>' if button_text and button_link else ''}
            </div>
        </div>
        """

        with tab2:
            st.write("### Inbox Preview")
            row = _preview_row(key_prefix + "_b", df, token)
            preview_html = replace_variables(generated_html, row)
            if local_inline_img:
                preview_html = preview_html.replace("cid:inline_img", "https://via.placeholder.com/600x150.png?text=Your+Uploaded+Banner+Will+Appear+Here")
            components.html(preview_html, height=500, scrolling=True)

    else:
        st.info(f"💡 **Pro Tip:** You can still use `{token}` and `{{Certificate ID}}` in your custom HTML. To embed a local image, upload it below and use `<img src=\"cid:inline_img\">` in your code.")

        local_inline_img = st.file_uploader(
            "Optional: Upload an image to embed as 'cid:inline_img'", type=['png', 'jpg', 'jpeg'], key=f"{key_prefix}_html_img"
        )

        html_tab, preview_tab = st.tabs(["💻 Paste HTML", "👁️ Live Preview"])
        with html_tab:
            generated_html = st.text_area(
                "Paste your raw HTML code here:", height=400,
                value=f"<h1>Hello {token}</h1>\n<p>Your custom HTML goes here.</p>",
                key=f"{key_prefix}_html_textarea",
            )
        with preview_tab:
            st.write("### Inbox Preview")
            row = _preview_row(key_prefix + "_h", df, token)
            preview_html = replace_variables(generated_html, row)
            if local_inline_img:
                preview_html = preview_html.replace("cid:inline_img", "https://via.placeholder.com/600x150.png?text=Your+Uploaded+Banner+Will+Appear+Here")
            components.html(preview_html, height=500, scrolling=True)

    return generated_html, local_inline_img


def render_certificate_section(key_prefix, df, csv_columns):
    attach_certificate = st.checkbox("📜 Generate & attach a personalized certificate to each email", key=f"{key_prefix}_attach_cert")
    cert_template_bytes = None
    placeholder_sources = {}
    cert_format = "PDF"

    if attach_certificate:
        cert_file = st.file_uploader("Upload certificate template (.pptx)", type=["pptx"], key=f"{key_prefix}_cert_file")
        if cert_file:
            cert_template_bytes = cert_file.read()
            try:
                placeholders = extract_placeholders(cert_template_bytes)
            except Exception as e:
                st.error(f"Could not read the certificate template: {e}")
                placeholders = []

            if not placeholders:
                st.warning("No `{{placeholder}}` text found in this template. Make sure your slide text uses double curly braces, e.g. `{{Name}}`.")
            else:
                st.success(f"Found {len(placeholders)} placeholder(s): " + ", ".join(placeholders))
                st.caption("Map each placeholder to a CSV column (personalized per recipient) or a fixed value used on every certificate.")
                for ph in placeholders:
                    clean_name = ph.strip("{} ").strip()
                    c1, c2, c3 = st.columns([2, 1, 2])
                    with c1:
                        st.markdown(f"**`{ph}`**")
                    with c2:
                        source = st.radio("Source", ["CSV column", "Fixed value"], key=f"{key_prefix}_src_{ph}", label_visibility="collapsed")
                    with c3:
                        if source == "CSV column":
                            if csv_columns:
                                default_idx = next(
                                    (i for i, c in enumerate(csv_columns) if c.strip().lower() == clean_name.lower()), 0
                                )
                                chosen_col = st.selectbox("Column", csv_columns, index=default_idx, key=f"{key_prefix}_col_{ph}", label_visibility="collapsed")
                            else:
                                chosen_col = None
                                st.caption("⬆️ Upload your audience CSV above to map this.")
                            placeholder_sources[ph] = ("column", chosen_col)
                        else:
                            fixed_val = st.text_input("Value", key=f"{key_prefix}_fixed_{ph}", label_visibility="collapsed")
                            placeholder_sources[ph] = ("fixed", fixed_val)

                format_choice = st.radio(
                    "Attach certificate as", ["PDF (recommended)", "Image (PNG)", "PPTX"], horizontal=True, key=f"{key_prefix}_cert_format"
                )
                if "PDF" in format_choice: cert_format = "PDF"
                elif "Image" in format_choice: cert_format = "Image"
                else: cert_format = "PPTX"

                if df is not None and len(df) > 0:
                    if st.button("🔍 Generate a sample certificate (first row)", key=f"{key_prefix}_sample_btn"):
                        sample_row = df.iloc[0].to_dict()
                        mapping = {}
                        for ph, (src_type, src_val) in placeholder_sources.items():
                            mapping[ph] = sample_row.get(src_val, "") if (src_type == "column" and src_val) else (src_val or "")
                        try:
                            filled = fill_certificate(cert_template_bytes, mapping)
                            if cert_format == "PDF":
                                pdf_bytes = pptx_bytes_to_pdf_bytes(filled)
                                if pdf_bytes is None:
                                    st.warning("⚠️ Couldn't convert to PDF. Showing .pptx instead.")
                                    st.download_button("Download sample (.pptx)", filled, file_name="sample_certificate.pptx", key=f"{key_prefix}_dl_pptx")
                                else:
                                    st.download_button("Download sample (.pdf)", pdf_bytes, file_name="sample_certificate.pdf", key=f"{key_prefix}_dl_pdf")
                            elif cert_format == "Image":
                                img_bytes = pptx_bytes_to_image_bytes(filled)
                                if img_bytes is None:
                                    st.warning("⚠️ Couldn't convert to Image. Showing .pptx instead.")
                                    st.download_button("Download sample (.pptx)", filled, file_name="sample_certificate.pptx", key=f"{key_prefix}_dl_pptx")
                                else:
                                    st.download_button("Download sample (.png)", img_bytes, file_name="sample_certificate.png", mime="image/png", key=f"{key_prefix}_dl_img")
                            else:
                                st.download_button("Download sample (.pptx)", filled, file_name="sample_certificate.pptx", key=f"{key_prefix}_dl_pptx2")
                        except Exception as e:
                            st.error(f"Couldn't generate sample certificate: {e}")
                else:
                    st.caption("Upload your audience CSV above to preview a sample certificate.")

    return attach_certificate, cert_template_bytes, placeholder_sources, cert_format


def render_subject_attachments(key_prefix, default_subject="Important Update"):
    col1, col2 = st.columns(2)
    with col1:
        subject = st.text_input("Email Subject", value=default_subject, key=f"{key_prefix}_subject",
                                help="You can use variables here too, e.g. {Name}")
    with col2:
        uploaded_attachments = st.file_uploader(
            "Attach files sent to everyone (PDF, Excel, etc.)", accept_multiple_files=True, key=f"{key_prefix}_attachments"
        )
    return subject, uploaded_attachments


def dispatch_campaign(key_prefix, df, generated_html, local_inline_img, subject, uploaded_attachments,
                       attach_certificate, cert_template_bytes, placeholder_sources, cert_format,
                       sender_name, sender_email, sender_password, name_col="Name", token="{Name}",
                       show_log=False):
    smtp_cfg = st.session_state.get("smtp_cfg", {})
    if df is None:
        st.info("Add your audience above to enable sending.")
        return

    st.warning("⚠️ Double-check your preview and data. Once you click send, the emails will be dispatched immediately.")
    if st.button("🚀 Confirm & Send Campaign", type="primary", key=f"{key_prefix}_send_btn"):
        if not all([sender_email, sender_password, subject]):
            st.error("❌ Please ensure your sender credentials and Subject line are filled out.")
            return
        if attach_certificate and not cert_template_bytes:
            st.error("❌ Certificate attachment is enabled but no template was uploaded.")
            return

        rows = df.to_dict('records')
        progress_bar = st.progress(0)
        status_text = st.empty()
        log_container = st.container()
        cert_fallback_warned = False
        log_rows = []

        try:
            status_text.text(f"Connecting to {smtp_cfg.get('host', 'smtp.zoho.in')}...")
            server = smtplib.SMTP(smtp_cfg.get("host", "smtp.zoho.in"), 587)
            server.starttls()
            server.login(smtp_cfg.get("user", sender_email), sender_password)
            status_text.text("✅ Authentication successful! Dispatching emails...")

            total_emails = len(rows)

            for index, row in enumerate(rows):
                recipient_name = row.get(name_col, 'Recipient')
                recipient_email = row.get('Email')
                cert_id = row.get('Certificate ID', '')
                send_ts = datetime.now().isoformat(timespec="seconds")

                if pd.isna(recipient_email) or not str(recipient_email).strip():
                    log_container.write(f"⚠️ Skipped {recipient_name}: No valid email address.")
                    log_rows.append({"Timestamp": send_ts, "Email": recipient_email, "Name": recipient_name,
                                      "Certificate ID": cert_id, "Status": "Skipped: no valid email"})
                    continue

                try:
                    personalized_body = replace_variables(generated_html, row)
                    personalized_subject = replace_variables(subject, row)

                    msg = MIMEMultipart('related')
                    msg['From'] = formataddr((sender_name, sender_email))
                    msg['To'] = str(recipient_email)
                    msg['Subject'] = personalized_subject

                    msg_alternative = MIMEMultipart('alternative')
                    msg.attach(msg_alternative)
                    msg_alternative.attach(MIMEText(personalized_body, 'html'))

                    if local_inline_img:
                        local_inline_img.seek(0)
                        image = MIMEImage(local_inline_img.read(), name=local_inline_img.name)
                        image.add_header('Content-ID', '<inline_img>')
                        image.add_header('Content-Disposition', 'inline')
                        msg.attach(image)

                    if uploaded_attachments:
                        for attached_file in uploaded_attachments:
                            attached_file.seek(0)
                            part = MIMEBase('application', 'octet-stream')
                            part.set_payload(attached_file.read())
                            encoders.encode_base64(part)
                            part.add_header('Content-Disposition', f'attachment; filename="{attached_file.name}"')
                            msg.attach(part)

                    if attach_certificate and cert_template_bytes:
                        mapping = {}
                        for ph, (src_type, src_val) in placeholder_sources.items():
                            mapping[ph] = row.get(src_val, "") if (src_type == "column" and src_val) else (src_val or "")

                        filled_pptx = fill_certificate(cert_template_bytes, mapping)
                        fname_base = sanitize_filename(recipient_name)
                        cert_bytes = filled_pptx
                        cert_filename = f"Certificate_{fname_base}.pptx"

                        if cert_format == "PDF":
                            pdf_bytes = pptx_bytes_to_pdf_bytes(filled_pptx)
                            if pdf_bytes is not None:
                                cert_bytes = pdf_bytes
                                cert_filename = f"Certificate_{fname_base}.pdf"
                            elif not cert_fallback_warned:
                                log_container.write("⚠️ PDF conversion unavailable — attaching .pptx certificates instead.")
                                cert_fallback_warned = True
                                
                        elif cert_format == "Image":
                            img_bytes = pptx_bytes_to_image_bytes(filled_pptx)
                            if img_bytes is not None:
                                cert_bytes = img_bytes
                                cert_filename = f"Certificate_{fname_base}.png"
                            elif not cert_fallback_warned:
                                log_container.write("⚠️ Image conversion failed — attaching .pptx certificates instead.")
                                cert_fallback_warned = True

                        cert_part = MIMEBase('application', 'octet-stream')
                        cert_part.set_payload(cert_bytes)
                        encoders.encode_base64(cert_part)
                        cert_part.add_header('Content-Disposition', f'attachment; filename="{cert_filename}"')
                        msg.attach(cert_part)

                    server.send_message(msg)
                    log_container.write(f"✅ Dispatched to: {recipient_name} ({recipient_email})")
                    log_rows.append({"Timestamp": send_ts, "Email": recipient_email, "Name": recipient_name,
                                      "Certificate ID": cert_id, "Status": "Sent"})

                except Exception as row_err:
                    log_container.write(f"❌ Failed for {recipient_name} ({recipient_email}): {row_err}")
                    log_rows.append({"Timestamp": send_ts, "Email": recipient_email, "Name": recipient_name,
                                      "Certificate ID": cert_id, "Status": f"Failed: {row_err}"})

                progress_bar.progress((index + 1) / total_emails)
                time.sleep(2)

            server.quit()
            status_text.success("🎉 Campaign completely dispatched!")

        except smtplib.SMTPAuthenticationError:
            st.error("❌ Authentication Error: Invalid sender email or password / API key.")
        except Exception as e:
            st.error(f"❌ A system error occurred: {e}")

        if show_log and log_rows:
            log_df = pd.DataFrame(log_rows, columns=["Timestamp", "Email", "Name", "Certificate ID", "Status"])
            st.session_state[f"{key_prefix}_log_df"] = log_df

    if show_log and st.session_state.get(f"{key_prefix}_log_df") is not None:
        log_df = st.session_state[f"{key_prefix}_log_df"]
        st.subheader("📋 Send Log")
        st.dataframe(log_df, use_container_width=True)
        st.download_button(
            "⬇️ Download log (CSV)",
            log_df.to_csv(index=False).encode("utf-8"),
            file_name="certificate_campaign_log.csv",
            mime="text/csv",
            key=f"{key_prefix}_log_download",
        )


# =========================================================================
# SIDEBAR — shared sender credentials
# =========================================================================
with st.sidebar:
    st.header("Server Credentials")
    provider = st.radio("Email service", ["ZeptoMail", "Zoho Mail"], horizontal=True)
    sender_name = st.text_input("Sender Display Name", value="Ei Updates")
    if provider == "ZeptoMail":
        st.info("💡 Use your ZeptoMail **Send Mail token**. The From address must be on a domain verified in ZeptoMail.")
        sender_email = st.text_input("From Email Address", value="@ei.study")
        raw_key = st.text_input("ZeptoMail API Key / SMTP Password", type="password")
        sender_password = re.sub(r"^\s*Zoho-enczapikey\s+", "", raw_key or "").strip()
        st.session_state["smtp_cfg"] = {"host": "smtp.zeptomail.in", "user": "emailapikey"}
    else:
        st.info("💡 **Security Note:** Never use your primary login password. Generate an App-Specific Password in your Zoho Security Settings.")
        sender_email = st.text_input("Zoho Email Address", value="@ei.study")
        sender_password = st.text_input("Zoho App Password", type="password")
        st.session_state["smtp_cfg"] = {"host": "smtp.zoho.in", "user": sender_email}
    st.caption("Used for sending both campaign types below.")


# =========================================================================
# TWO INDEPENDENT CAMPAIGNS
# =========================================================================
campaign_tab1, campaign_tab2 = st.tabs(["📧 Regular Email Campaign", "🎓 Certificate Campaign"])

with campaign_tab1:
    st.header("1. Audience")
    df_reg, cols_reg = render_audience("reg", name_col="Name")

    st.write("---")
    st.header("2. Design Your Email")
    html_reg, img_reg = render_compose(
        "reg", 
        default_heading="Welcome to {Event}, {School Name}!", 
        token="{Name}", 
        df=df_reg
    )

    st.write("---")
    st.header("3. Subject & Attachments")
    subject_reg, attachments_reg = render_subject_attachments(
        "reg", 
        default_subject="An Invitation to Connect Regarding {Event}"
    )

    st.write("---")
    st.header("4. Dispatch")
    dispatch_campaign(
        "reg", df_reg, html_reg, img_reg, subject_reg, attachments_reg,
        attach_certificate=False, cert_template_bytes=None, placeholder_sources={}, cert_format="PDF",
        sender_name=sender_name, sender_email=sender_email, sender_password=sender_password,
        name_col="Name", token="{Name}", show_log=False,
    )

with campaign_tab2:
    st.header("1. Audience")
    st.caption(
        "Independent of the Regular Campaign tab — upload your response sheet "
        "(e.g. `Timestamp, Email, Name`, exported from a Google Form) just for certificate recipients."
    )
    df_cert, cols_cert = render_audience("cert", name_col="Name")

    if df_cert is not None and len(df_cert) > 0:
        st.write("")
        id_prefix = st.text_input("Certificate ID prefix", value="EI-CERT-", key="cert_id_prefix",
                                   help="Certificate IDs are auto-generated as <prefix><sequence>, e.g. EI-CERT-00001")
        df_cert = df_cert.reset_index(drop=True)
        df_cert["Certificate ID"] = [f"{id_prefix}{i+1:05d}" for i in range(len(df_cert))]
        cols_cert = list(df_cert.columns)
        with st.expander("🪪 Certificate IDs assigned", expanded=False):
            display_cols = [c for c in ["Email", "Name", "Certificate ID"] if c in df_cert.columns]
            st.dataframe(df_cert[display_cols] if display_cols else df_cert, use_container_width=True)

    st.write("---")
    st.header("2. Design Your Certificate Email")
    html_cert, img_cert = render_compose(
        "cert", default_heading="Your Certificate is Here, {Name}!", token="{Name}", df=df_cert
    )

    st.write("---")
    st.header("3. Certificate")
    attach_cert, cert_bytes, cert_sources, cert_format = render_certificate_section("cert", df_cert, cols_cert)

    st.write("---")
    st.header("4. Subject & Attachments")
    subject_cert, attachments_cert = render_subject_attachments("cert", default_subject="Your Certificate of Participation")

    st.write("---")
    st.header("5. Dispatch")
    dispatch_campaign(
        "cert", df_cert, html_cert, img_cert, subject_cert, attachments_cert,
        attach_certificate=attach_cert, cert_template_bytes=cert_bytes,
        placeholder_sources=cert_sources, cert_format=cert_format,
        sender_name=sender_name, sender_email=sender_email, sender_password=sender_password,
        name_col="Name", token="{Name}", show_log=True,
    )